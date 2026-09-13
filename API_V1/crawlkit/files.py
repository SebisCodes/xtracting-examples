"""Files on a page, and how they become text.

XTRACTING TAKES TEXT, NOTHING ELSE. So a PDF is not "attached" anywhere - its
text layer is extracted here and submitted like a page. What has no text layer
(a scan, a drawing, an empty sheet) is not submitted at all: an empty task is
billed and yields nothing.

THREE RULES, EACH FROM AN INCIDENT:

  1. WHAT ARRIVES IS CHECKED, NOT WHAT THE ADDRESS PROMISED. An address ending
     in `.pdf` can answer with an HTML error page, and an address without an
     extension can be a PDF. The extension is only the reason to look; the
     bytes decide.

  2. THE SIZE IS CHECKED WHILE LOADING, NOT AFTERWARDS. A 30 MB plan set must
     not sit in memory before anyone notices it is too big.

  3. WHAT WAS NOT SENT LEAVES A ROW ANYWAY, with its reason. "Why is that PDF
     not in the archive?" is the question `scraper.files` exists to answer,
     and "too large", "no text layer" and "the type is switched off" are three
     different answers.
"""

from __future__ import annotations

import csv as csv_module
import io
import logging
import re
import zipfile
from contextlib import contextmanager
from dataclasses import dataclass, field
from typing import Iterator
from urllib.parse import unquote, urlsplit

from crawlkit.classify import FILE_EXTENSIONS
from crawlkit.content import extract_links
from crawlkit.hashing import canonical_uri

log = logging.getLogger(__name__)

#: The types that can become text. Everything else in `FILE_EXTENSIONS` is
#: recognised as a file and then listed as "cannot be sent as text".
TEXT_TYPES = ("pdf", "docx", "xlsx", "csv", "txt", "md")

#: The first bytes of the formats we accept. The last check when content type
#: and extension disagree - and they disagree exactly when it matters.
MAGIC = {"pdf": b"%PDF-", "docx": b"PK\x03\x04", "xlsx": b"PK\x03\x04"}

#: Read in this size while checking the cap. Small enough that the cap is not
#: overshot by much, large enough not to make a syscall per kilobyte.
CHUNK_BYTES = 64 * 1024

#: The statuses of `scraper.files`, mirrored here so that this module can name
#: them without importing the database layer.
STATUSES = ("QUEUED", "SENT", "TOO_LARGE", "WRONG_TYPE", "NO_TEXT", "ROBOTS",
            "ERROR", "NOT_SELECTED")


def file_type(url: str) -> str:
    """The type of a file address - lower case, without the dot.

    >>> file_type("https://a.example/files/expose-1.PDF?download=1")
    'pdf'
    >>> file_type("https://a.example/page")
    ''
    """
    last = urlsplit(url).path.rsplit("/", 1)[-1]
    ext = last.rsplit(".", 1)[-1].lower() if "." in last else ""
    return ext if ext in FILE_EXTENSIONS else ""


def file_name(url: str) -> str:
    """The file name from the address, for display.

    >>> file_name("https://a.example/files/Site%20plan%20A.pdf?x=1")
    'Site plan A.pdf'
    >>> file_name("https://a.example/download?id=1")
    'download'
    """
    path = urlsplit(url).path.rstrip("/")
    name = unquote(path.rsplit("/", 1)[-1]) if path else ""
    return name[:200] or "file"


@dataclass(frozen=True)
class FoundFile:
    url: str
    type: str
    from_page: str = ""
    text: str = ""


def discover(html: str, page_url: str) -> list[FoundFile]:
    """The files linked from one page.

    NOT limited to the page's own host, and that is the one place where this
    is right: portals regularly serve their documents from a separate file
    server or a cantonal portal, and the host brake would leave out exactly
    the documents that were the point. It is not dangerous either - only what
    looks like a file is looked at, and only what proves to be one is read.

    >>> html = ('<a href="/f/expose-1.pdf">Expose</a>'
    ...         '<a href="https://cdn.example/plan.docx">Plan</a>'
    ...         '<a href="/rent/2">another flat</a>')
    >>> [(f.type, f.url) for f in discover(html, "https://a.example/rent/1")]
    [('pdf', 'https://a.example/f/expose-1.pdf'), ('docx', 'https://cdn.example/plan.docx')]
    """
    out: list[FoundFile] = []
    for link in extract_links(html, page_url):
        kind = file_type(link.url)
        if kind:
            out.append(FoundFile(link.url, kind, page_url, link.text))
    return out


@dataclass
class FileRules:
    """Which files a source sends: a default per type, exceptions per address.

    The two are different concepts and the popup keeps them visibly apart -
    "send PDFs" is a decision about a hundred files, "not this one" about one.

    >>> rules = FileRules.from_config({"file_rules": [
    ...     {"text_kind": "type_default", "text_value": "pdf", "bool_enabled": True,
    ...      "integer_max_mb": 25},
    ...     {"text_kind": "type_default", "text_value": "xlsx", "bool_enabled": False},
    ...     {"text_kind": "file_exclude", "text_value": "https://a.example/f/big.pdf"},
    ... ]})
    >>> rules.wanted("https://a.example/f/one.pdf", "pdf")
    (True, '')
    >>> rules.wanted("https://a.example/f/big.pdf", "pdf")
    (False, 'this file is excluded')
    >>> rules.wanted("https://a.example/f/book.xlsx", "xlsx")
    (False, 'xlsx files are switched off for this watched page')
    >>> rules.wanted("https://a.example/f/x.zip", "zip")
    (False, 'zip cannot be sent as text')
    """

    types: dict[str, bool] = field(default_factory=dict)
    include: set[str] = field(default_factory=set)
    exclude: set[str] = field(default_factory=set)
    max_mb: dict[str, int] = field(default_factory=dict)
    default_max_mb: int = 25

    @staticmethod
    def from_config(config: dict) -> "FileRules":
        rules = FileRules()
        for row in config.get("file_rules") or ():
            kind = str(row.get("text_kind") or "")
            value = str(row.get("text_value") or "").strip()
            if not value:
                continue
            if kind == "type_default":
                rules.types[value.lower()] = bool(row.get("bool_enabled", True))
                rules.max_mb[value.lower()] = int(row.get("integer_max_mb") or 25)
            elif kind == "file_include":
                rules.include.add(canonical_uri(value))
            elif kind == "file_exclude":
                rules.exclude.add(canonical_uri(value))
        return rules

    def cap_bytes(self, kind: str) -> int:
        return int(self.max_mb.get(kind, self.default_max_mb)) * 1024 * 1024

    def wanted(self, url: str, kind: str) -> tuple[bool, str]:
        """Send this file? The reason is what the file list shows."""
        canon = canonical_uri(url)
        if canon in self.exclude:
            return False, "this file is excluded"
        if canon in self.include:
            return True, ""
        if kind not in TEXT_TYPES:
            return False, f"{kind or 'this file'} cannot be sent as text"
        if not self.types.get(kind, False):
            return False, f"{kind} files are switched off for this watched page"
        return True, ""


@dataclass
class Download:
    """What came of one attempt. `status` is a `scraper.files` status."""

    url: str
    status: str
    detail: str = ""
    data: bytes = b""
    content_type: str = ""
    bytes_read: int = 0

    @property
    def ok(self) -> bool:
        return self.status == "QUEUED"


@contextmanager
def _open_stream(url: str, *, user_agent: str, timeout: float):
    import httpx
    with httpx.stream("GET", url, timeout=timeout, follow_redirects=True,
                      headers={"User-Agent": user_agent,
                               "Accept": "application/pdf,*/*;q=0.8"}) as response:
        yield response


def download(url: str, kind: str, *, user_agent: str, max_bytes: int,
             timeout: float = 60.0, opener=_open_stream) -> Download:
    """Load one file, with a cap and a type check.

    `opener` is injectable so the tests can hand over a stream that never
    ends: what is being checked is that the reading stops, and it must be
    checkable without a 30 MB download.
    """
    try:
        with opener(url, user_agent=user_agent, timeout=timeout) as response:
            status = getattr(response, "status_code", 0)
            headers = getattr(response, "headers", {}) or {}
            ctype = str(headers.get("content-type", "")).split(";")[0].strip()

            if status >= 400:
                return Download(url, "ERROR", f"HTTP {status}", content_type=ctype)

            # The server names the size before anything is loaded. If it is
            # honest, that saves the whole download.
            announced = str(headers.get("content-length", "")).strip()
            if announced.isdigit() and int(announced) > max_bytes:
                return Download(url, "TOO_LARGE",
                                f"{int(announced) // 1024 // 1024} MB, the limit "
                                f"is {max_bytes // 1024 // 1024} MB",
                                content_type=ctype)

            chunks: list[bytes] = []
            read = 0
            for chunk in response.iter_bytes(CHUNK_BYTES):
                read += len(chunk)
                if read > max_bytes:
                    # Stop reading, do not finish the download. Otherwise the
                    # cap is bookkeeping about bytes that already came down
                    # the wire.
                    return Download(url, "TOO_LARGE",
                                    f"more than {max_bytes // 1024 // 1024} MB "
                                    f"(the server did not announce the size)",
                                    content_type=ctype, bytes_read=read)
                chunks.append(chunk)
    except Exception as exc:  # noqa: BLE001 - every cause ends the same way
        return Download(url, "ERROR", f"{type(exc).__name__}: {exc}"[:400])

    data = b"".join(chunks)
    if not data:
        return Download(url, "ERROR", "empty answer", content_type=ctype, bytes_read=0)

    magic = MAGIC.get(kind)
    if magic and not data.startswith(magic):
        # THE CHECK THAT MATTERS. A login mask served instead of the document
        # answers 200 with text/html - without this line an HTML error page
        # would be archived as a PDF.
        return Download(url, "WRONG_TYPE",
                        f"not a {kind} (content type {ctype or 'unknown'}, "
                        f"starts with {data[:8]!r})",
                        content_type=ctype, bytes_read=len(data))
    if kind == "xlsx" and not _is_ooxml(data):
        return Download(url, "WRONG_TYPE",
                        f"not an xlsx (a zip without [Content_Types].xml)",
                        content_type=ctype, bytes_read=len(data))

    return Download(url, "QUEUED", "", data=data, content_type=ctype,
                    bytes_read=len(data))


def _is_ooxml(data: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            return "[Content_Types].xml" in archive.namelist()
    except (zipfile.BadZipFile, OSError):
        return False


def to_text(data: bytes, kind: str) -> tuple[str, str]:
    """(text, reason). An empty text always comes with a reason.

    >>> to_text(b"Name;Rent\\nFlat;1200\\n", "csv")[0]
    'Name;Rent\\nFlat;1200'
    >>> to_text(b"", "txt")
    ('', 'the file is empty')

    A file that cannot be read is a reason, never an exception: one broken
    document must not end a run that still has twenty good ones in it.

    >>> to_text(b"PK\x03\x04 truncated", "docx")[0]
    ''
    >>> to_text(b"PK\x03\x04 truncated", "docx")[1].startswith("the file could not be read")
    True

    A PDF with pages and no text layer - a scan - is the case where the file
    is fine and there is still nothing to send. `test_files.py` builds one and
    checks the wording, because that wording is what the file list shows.
    """
    if not data:
        return "", "the file is empty"
    try:
        if kind == "pdf":
            return _pdf_text(data)
        if kind == "docx":
            return _docx_text(data)
        if kind == "xlsx":
            return _xlsx_text(data)
        if kind in ("csv", "txt", "md"):
            text = _decode(data)
            return (text.strip(), "") if text.strip() else ("", "the file is empty")
    except Exception as exc:  # noqa: BLE001 - a broken file is not a broken run
        return "", f"the file could not be read ({type(exc).__name__}: {exc})"[:400]
    return "", f"{kind} cannot be sent as text"


def _decode(data: bytes) -> str:
    """Bytes to text, trying the encodings that actually turn up.

    A CSV exported from a spreadsheet on Windows is cp1252 often enough that
    failing on it would mean losing the file over an umlaut.

    >>> _decode("Preis;Größe\\n".encode("cp1252"))
    'Preis;Größe\\n'
    """
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf_text(data: bytes) -> tuple[str, str]:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data), strict=False)
    parts = []
    for page in reader.pages:
        parts.append((page.extract_text() or "").strip())
    text = "\n\n".join(part for part in parts if part).strip()
    if not text:
        return "", "the PDF has no text layer - a scan or a drawing"
    return text, ""


def _docx_text(data: bytes) -> tuple[str, str]:
    """The paragraphs of a .docx.

    Read straight from the zip rather than with a library: a .docx is a zip
    with an XML in it, the paragraphs are `<w:p>` elements, and one dependency
    fewer in an image is one dependency fewer to keep current.
    """
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if "word/document.xml" not in archive.namelist():
            return "", "not a Word document (word/document.xml is missing)"
        xml = archive.read("word/document.xml").decode("utf-8", "replace")
    xml = re.sub(r"</w:p>", "\n", xml)
    xml = re.sub(r"<w:tab[^>]*/>", "\t", xml)
    text = re.sub(r"<[^>]+>", "", xml)
    text = re.sub(r"[ \t]+\n", "\n", text).strip()
    if not text:
        return "", "the document has no text"
    return text, ""


def _xlsx_text(data: bytes) -> tuple[str, str]:
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for sheet in book.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if cell is None else str(cell) for cell in row]
                if any(cell.strip() for cell in cells):
                    rows.append("\t".join(cells).rstrip())
            if rows:
                # The sheet name is content: "2024" or "Costs" says what the
                # numbers below it are.
                parts.append(f"{sheet.title}\n" + "\n".join(rows))
    finally:
        book.close()
    text = "\n\n".join(parts).strip()
    if not text:
        return "", "the workbook has no cells with content"
    return text, ""


def rows_of_csv(text: str) -> Iterator[list[str]]:
    """The rows of a CSV, for whoever wants them - the submission takes the
    text as it is."""
    return csv_module.reader(io.StringIO(text))


__all__ = ["CHUNK_BYTES", "Download", "FileRules", "FoundFile", "MAGIC",
           "STATUSES", "TEXT_TYPES", "discover", "download", "file_name",
           "file_type", "rows_of_csv", "to_text"]
